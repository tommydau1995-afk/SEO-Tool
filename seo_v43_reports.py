"""PDF and Excel report exporters for SEO AI Studio V4.3."""

from __future__ import annotations

import os
from collections import Counter
from dataclasses import asdict, is_dataclass
from datetime import datetime
from pathlib import Path

from seo_v43_core import VERSION, dashboard_kpis


NAVY = "17324D"
BLUE = "1769C2"
TEAL = "0E8A7A"
GREEN = "2E7D32"
AMBER = "F59E0B"
RED = "C62828"
PALE_BLUE = "EAF3FF"
PALE_GREEN = "EAF7F2"
PALE_RED = "FDECEC"
LIGHT = "F5F7FA"
WHITE = "FFFFFF"
GRAY = "64748B"


def _row(value):
    if is_dataclass(value):
        return asdict(value)
    return dict(value)


def report_payload(**kwargs):
    """Normalize report input to keep the GUI and scheduled runner decoupled."""

    payload = {
        "project": kwargs.get("project") or {},
        "pages": [_row(value) for value in kwargs.get("pages", [])],
        "fixes": [_row(value) for value in kwargs.get("fixes", [])],
        "site_files": kwargs.get("site_files") or {},
        "link_map": kwargs.get("link_map") or {},
        "log_report": kwargs.get("log_report") or {},
        "gsc_rows": list(kwargs.get("gsc_rows", [])),
        "pagespeed": kwargs.get("pagespeed") or {},
        "keyword_rows": list(kwargs.get("keyword_rows", [])),
        "backlinks": list(kwargs.get("backlinks", [])),
        "backlink_summary": kwargs.get("backlink_summary") or {},
        "decay_rows": list(kwargs.get("decay_rows", [])),
        "rank_rows": list(kwargs.get("rank_rows", [])),
        "roadmap": list(kwargs.get("roadmap", [])),
        "master_checks": [_row(value) for value in kwargs.get("master_checks", [])],
        "alerts": kwargs.get("alerts") or {},
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "version": VERSION,
    }
    payload["kpis"] = dashboard_kpis(
        kwargs.get("pages", []),
        kwargs.get("fixes", []),
        payload["gsc_rows"],
        payload["log_report"],
        payload["decay_rows"],
        payload["backlinks"],
    )
    return payload


def export_excel(path, payload):
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import CellIsRule, FormulaRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError(
            "Thiếu openpyxl. Hãy cài lại SEO AI Studio V4.3 bằng bộ cài đầy đủ."
        ) from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_view.showGridLines = False

    thin = Side(style="thin", color="DDE3EA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_fill = PatternFill("solid", fgColor=NAVY)
    header_fill = PatternFill("solid", fgColor=BLUE)

    ws.merge_cells("A1:H2")
    ws["A1"] = "SEO AI STUDIO V4.3 - BÁO CÁO TỔNG QUAN"
    ws["A1"].font = Font(size=20, bold=True, color=WHITE)
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left")
    project = payload.get("project", {})
    ws["A4"] = "Dự án"
    ws["B4"] = project.get("name", "Chưa đặt tên")
    ws["A5"] = "Website"
    ws["B5"] = project.get("website", "")
    ws["A6"] = "Tạo lúc"
    ws["B6"] = payload.get("generated_at", "")
    for cell in ("A4", "A5", "A6"):
        ws[cell].font = Font(bold=True, color=NAVY)

    cards = [
        ("SEO Score", payload["kpis"]["seo_score"], "/100"),
        ("Trang Audit", payload["kpis"]["pages"], ""),
        ("P0 đang mở", payload["kpis"]["p0_open"], ""),
        ("GSC Clicks", payload["kpis"]["gsc_clicks"], ""),
        ("Crawl hiệu quả", payload["kpis"]["crawl_efficiency"], "%"),
        ("Content Decay", payload["kpis"]["decay_pages"], " trang"),
    ]
    for index, (label, value, suffix) in enumerate(cards):
        row = 8 + (index // 3) * 4
        col = 1 + (index % 3) * 3
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.merge_cells(
            start_row=row + 1, start_column=col, end_row=row + 2, end_column=col + 1
        )
        label_cell = ws.cell(row, col, label)
        value_cell = ws.cell(row + 1, col, f"{value}{suffix}")
        label_cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
        label_cell.font = Font(bold=True, color=BLUE)
        label_cell.alignment = Alignment(horizontal="center")
        value_cell.fill = PatternFill("solid", fgColor=WHITE)
        value_cell.font = Font(size=18, bold=True, color=NAVY)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        for r in range(row, row + 3):
            for c in range(col, col + 2):
                ws.cell(r, c).border = border

    ws["A17"] = "Tình trạng ưu tiên"
    ws["A17"].font = Font(size=14, bold=True, color=NAVY)
    ws["A18"] = "P0"
    ws["B18"] = '=COUNTIF(\'Fix Queue\'!A:A,"P0")'
    ws["A19"] = "P1"
    ws["B19"] = '=COUNTIF(\'Fix Queue\'!A:A,"P1")'
    ws["A20"] = "P2"
    ws["B20"] = '=COUNTIF(\'Fix Queue\'!A:A,"P2")'
    ws["D17"] = "Roadmap"
    ws["D17"].font = Font(size=14, bold=True, color=NAVY)
    ws["D18"] = "30 ngày"
    ws["E18"] = '=COUNTIF(Roadmap!A:A,30)'
    ws["D19"] = "60 ngày"
    ws["E19"] = '=COUNTIF(Roadmap!A:A,60)'
    ws["D20"] = "90 ngày"
    ws["E20"] = '=COUNTIF(Roadmap!A:A,90)'
    ws.freeze_panes = "A4"
    ws.column_dimensions["A"].width = 19
    ws.column_dimensions["B"].width = 28
    for col in ("C", "D", "E", "F", "G", "H"):
        ws.column_dimensions[col].width = 18

    def add_sheet(name, headers, rows, widths=None):
        sheet = wb.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        for column, header in enumerate(headers, 1):
            cell = sheet.cell(1, column, header)
            cell.fill = header_fill
            cell.font = Font(bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        for row_index, source in enumerate(rows, 2):
            source = _row(source)
            for column, key in enumerate(headers.keys(), 1):
                value = source.get(key, "")
                if isinstance(value, (list, dict)):
                    value = str(value)
                cell = sheet.cell(row_index, column, value)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 25
        for index, (key, label) in enumerate(headers.items(), 1):
            width = (widths or {}).get(key, min(55, max(12, len(label) + 4)))
            sheet.column_dimensions[get_column_letter(index)].width = width
        if rows:
            safe_name = re_table_name(name)
            table = Table(displayName=safe_name, ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)
        return sheet

    pages = []
    for page in payload["pages"]:
        row = dict(page)
        issues = row.get("issues", "")
        if isinstance(issues, list):
            row["issues"] = " | ".join(issues)
        pages.append(row)
    audit = add_sheet(
        "Audit",
        {
            "score": "Điểm",
            "status": "HTTP",
            "depth": "Depth",
            "inlinks": "Inlinks",
            "url": "URL",
            "title": "Title",
            "words": "Số từ",
            "response_ms": "Response ms",
            "indexable": "Indexable",
            "issues": "Lỗi",
        },
        pages,
        {"url": 55, "title": 45, "issues": 65},
    )
    if len(pages) >= 1:
        audit.conditional_formatting.add(
            f"A2:A{len(pages) + 1}",
            CellIsRule(
                operator="lessThan",
                formula=["60"],
                fill=PatternFill("solid", fgColor=PALE_RED),
            ),
        )

    fix_sheet = add_sheet(
        "Fix Queue",
        {
            "priority": "Ưu tiên",
            "issue": "Lỗi",
            "url": "URL",
            "recommendation": "Cách xử lý",
            "status": "Trạng thái",
        },
        payload["fixes"],
        {"issue": 28, "url": 55, "recommendation": 70, "status": 16},
    )
    if payload["fixes"]:
        end = len(payload["fixes"]) + 1
        fix_sheet.conditional_formatting.add(
            f"A2:A{end}",
            FormulaRule(formula=['A2="P0"'], fill=PatternFill("solid", fgColor="F8C7C7")),
        )
        fix_sheet.conditional_formatting.add(
            f"A2:A{end}",
            FormulaRule(formula=['A2="P1"'], fill=PatternFill("solid", fgColor="FFE8B0")),
        )

    add_sheet(
        "Internal Links",
        {
            "source": "Nguồn",
            "source_depth": "Depth nguồn",
            "target": "Đích",
            "target_depth": "Depth đích",
            "target_inlinks": "Inlinks đích",
            "target_status": "HTTP đích",
        },
        payload["link_map"].get("edges", []),
        {"source": 55, "target": 55},
    )
    add_sheet(
        "GSC 90 Days",
        {
            "date": "Ngày",
            "page": "Trang",
            "query": "Truy vấn",
            "clicks": "Clicks",
            "impressions": "Impressions",
            "ctr": "CTR",
            "position": "Position",
        },
        payload["gsc_rows"],
        {"page": 55, "query": 35},
    )
    add_sheet(
        "Content Decay",
        {
            "page": "Trang",
            "previous_clicks": "Clicks kỳ trước",
            "current_clicks": "Clicks kỳ này",
            "click_change_percent": "% thay đổi clicks",
            "previous_impressions": "Impressions kỳ trước",
            "current_impressions": "Impressions kỳ này",
            "impression_change_percent": "% thay đổi impressions",
            "status": "Phân loại",
        },
        payload["decay_rows"],
        {"page": 60, "status": 18},
    )
    add_sheet(
        "Backlinks",
        {
            "source": "Nguồn",
            "domain": "Referring domain",
            "target": "Đích",
            "anchor": "Anchor",
            "authority": "DR/DA",
            "traffic": "Traffic",
            "follow": "Dofollow",
            "risk": "Rủi ro",
        },
        payload["backlinks"],
        {"source": 55, "target": 55, "anchor": 30},
    )
    add_sheet(
        "Rank Tracking",
        {
            "checked_at": "Thời điểm",
            "keyword": "Từ khóa",
            "domain": "Domain",
            "position": "Vị trí",
            "link": "URL xếp hạng",
        },
        payload["rank_rows"],
        {"keyword": 35, "domain": 30, "link": 55},
    )
    add_sheet(
        "Keyword Map",
        {
            "keyword": "Từ khóa",
            "url": "Landing Page",
            "intent": "Intent",
            "priority": "Ưu tiên",
            "rank": "Rank",
            "status": "Trạng thái",
        },
        payload["keyword_rows"],
        {"keyword": 35, "url": 55},
    )
    add_sheet(
        "Roadmap",
        {
            "phase_days": "Giai đoạn (ngày)",
            "workstream": "Nhóm",
            "priority": "Ưu tiên",
            "action": "Hành động",
            "url": "URL",
            "success_metric": "Tiêu chí thành công",
            "status": "Trạng thái",
        },
        payload["roadmap"],
        {"action": 40, "url": 55, "success_metric": 55},
    )
    add_sheet(
        "SEO 150",
        {
            "id": "#",
            "category": "Nhóm",
            "priority": "Ưu tiên",
            "method": "Nguồn",
            "title": "Hạng mục",
            "status": "Trạng thái",
            "evidence": "Bằng chứng",
        },
        payload["master_checks"],
        {"category": 28, "title": 70, "evidence": 50},
    )

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(path)
    return path


def re_table_name(name):
    safe = "".join(character if character.isalnum() else "_" for character in name)
    if safe[:1].isdigit():
        safe = "T_" + safe
    return (safe or "Report")[:240]


def _font_paths():
    candidates = [
        (
            r"C:\Windows\Fonts\arial.ttf",
            r"C:\Windows\Fonts\arialbd.ttf",
        ),
        (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        ),
    ]
    for normal, bold in candidates:
        if Path(normal).exists() and Path(bold).exists():
            return normal, bold
    return None, None


def export_pdf(path, payload):
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            BaseDocTemplate,
            Frame,
            KeepTogether,
            PageBreak,
            PageTemplate,
            Paragraph,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:
        raise RuntimeError(
            "Thiếu reportlab. Hãy cài lại SEO AI Studio V4.3 bằng bộ cài đầy đủ."
        ) from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    normal_path, bold_path = _font_paths()
    font_name = "Helvetica"
    bold_name = "Helvetica-Bold"
    if normal_path and bold_path:
        pdfmetrics.registerFont(TTFont("SEOFont", normal_path))
        pdfmetrics.registerFont(TTFont("SEOFontBold", bold_path))
        font_name, bold_name = "SEOFont", "SEOFontBold"

    page_size = landscape(A4)
    page_width, page_height = page_size
    doc = BaseDocTemplate(
        str(path),
        pagesize=page_size,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=18 * mm,
        bottomMargin=15 * mm,
        title=f"SEO AI Studio V4.3 - {payload.get('project', {}).get('name', '')}",
        author="SEO AI Studio V4.3",
    )
    frame = Frame(
        doc.leftMargin,
        doc.bottomMargin,
        doc.width,
        doc.height,
        id="report",
    )

    def page_header_footer(canvas, document):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#" + NAVY))
        canvas.rect(0, page_height - 12 * mm, page_width, 12 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont(bold_name, 9)
        canvas.drawString(14 * mm, page_height - 8 * mm, "SEO AI STUDIO V4.3")
        canvas.setFillColor(colors.HexColor("#" + GRAY))
        canvas.setFont(font_name, 8)
        canvas.drawString(
            14 * mm,
            7 * mm,
            clean_pdf_text(payload.get("project", {}).get("website", "")),
        )
        canvas.drawRightString(
            page_width - 14 * mm,
            7 * mm,
            f"Trang {document.page}",
        )
        canvas.restoreState()

    doc.addPageTemplates(
        [PageTemplate(id="SEOReport", frames=[frame], onPage=page_header_footer)]
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="SEO_Title",
            fontName=bold_name,
            fontSize=24,
            leading=29,
            textColor=colors.HexColor("#" + NAVY),
            spaceAfter=5 * mm,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SEO_H1",
            fontName=bold_name,
            fontSize=15,
            leading=19,
            textColor=colors.HexColor("#" + NAVY),
            spaceBefore=3 * mm,
            spaceAfter=3 * mm,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SEO_Body",
            fontName=font_name,
            fontSize=8.5,
            leading=11.5,
            textColor=colors.HexColor("#" + NAVY),
        )
    )
    styles.add(
        ParagraphStyle(
            name="SEO_CardLabel",
            fontName=font_name,
            fontSize=8,
            leading=10,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#" + GRAY),
        )
    )
    styles.add(
        ParagraphStyle(
            name="SEO_CardValue",
            fontName=bold_name,
            fontSize=17,
            leading=20,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#" + BLUE),
        )
    )
    body = styles["SEO_Body"]
    story = []
    project = payload.get("project", {})
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("BÁO CÁO SEO 30/60/90 NGÀY", styles["SEO_Title"]))
    story.append(
        Paragraph(
            f"<b>{clean_pdf_text(project.get('name', 'Dự án SEO'))}</b><br/>"
            f"{clean_pdf_text(project.get('website', ''))}<br/>"
            f"Tạo lúc: {clean_pdf_text(payload.get('generated_at', ''))}",
            body,
        )
    )
    story.append(Spacer(1, 7 * mm))

    kpis = payload["kpis"]
    cards = [
        ("SEO Score", f"{kpis['seo_score']}/100"),
        ("Trang Audit", str(kpis["pages"])),
        ("P0 đang mở", str(kpis["p0_open"])),
        ("GSC Clicks", format_number(kpis["gsc_clicks"])),
        ("Crawl hiệu quả", f"{kpis['crawl_efficiency']}%"),
        ("Content Decay", str(kpis["decay_pages"])),
    ]
    card_data = []
    for label, value in cards:
        card_data.append(
            [
                Paragraph(value, styles["SEO_CardValue"]),
                Paragraph(label, styles["SEO_CardLabel"]),
            ]
        )
    card_table = Table(
        [card_data[:3], card_data[3:]],
        colWidths=[doc.width / 3] * 3,
        rowHeights=[23 * mm, 23 * mm],
    )
    card_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + PALE_BLUE)),
                ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCD8E5")),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCD8E5")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ]
        )
    )
    story.append(card_table)
    story.append(Spacer(1, 6 * mm))
    issue_counter = Counter(
        issue
        for page in payload["pages"]
        for issue in (
            page.get("issues", [])
            if isinstance(page.get("issues", []), list)
            else str(page.get("issues", "")).split(" | ")
        )
        if issue
    )
    issue_rows = [["Lỗi phổ biến", "Số trang"]] + [
        [name, count] for name, count in issue_counter.most_common(10)
    ]
    story.append(Paragraph("TỔNG QUAN VẤN ĐỀ", styles["SEO_H1"]))
    story.append(styled_table(issue_rows, [doc.width * 0.82, doc.width * 0.18], font_name, bold_name))

    story.append(PageBreak())
    story.append(Paragraph("FIX QUEUE ƯU TIÊN", styles["SEO_H1"]))
    fixes = sorted(
        payload["fixes"],
        key=lambda row: ({"P0": 0, "P1": 1, "P2": 2}.get(row.get("priority"), 9), row.get("url", "")),
    )[:35]
    fix_rows = [["Ưu tiên", "Lỗi", "URL", "Khuyến nghị", "Trạng thái"]]
    fix_rows.extend(
        [
            row.get("priority", ""),
            row.get("issue", ""),
            row.get("url", ""),
            row.get("recommendation", ""),
            row.get("status", ""),
        ]
        for row in fixes
    )
    story.append(
        styled_table(
            fix_rows,
            [18 * mm, 35 * mm, 68 * mm, 105 * mm, 24 * mm],
            font_name,
            bold_name,
            wrap=True,
        )
    )

    story.append(PageBreak())
    story.append(Paragraph("CRAWL, SITEMAP & INTERNAL LINKS", styles["SEO_H1"]))
    link_map = payload["link_map"]
    crawl_rows = [
        ["Chỉ số", "Giá trị"],
        ["Crawl efficiency", f"{payload['log_report'].get('crawl_efficiency', 0)}%"],
        ["Googlebot hits", payload["log_report"].get("googlebot_hits", 0)],
        ["Waste hits", payload["log_report"].get("waste_hits", 0)],
        ["Sitemap URLs", link_map.get("sitemap_urls", 0)],
        ["Sitemap coverage", f"{link_map.get('sitemap_coverage', 0)}%"],
        ["Potential orphan", len(link_map.get("orphans", []))],
        ["Dead-end", len(link_map.get("dead_ends", []))],
        ["Depth > 3", len(link_map.get("deep_pages", []))],
    ]
    story.append(styled_table(crawl_rows, [doc.width * 0.65, doc.width * 0.35], font_name, bold_name))
    story.append(Spacer(1, 5 * mm))
    story.append(Paragraph("URL cần chú ý", styles["SEO_H1"]))
    attention = (
        [("Orphan", value) for value in link_map.get("orphans", [])[:12]]
        + [("Dead-end", value) for value in link_map.get("dead_ends", [])[:10]]
        + [("Depth > 3", value) for value in link_map.get("deep_pages", [])[:10]]
    )
    story.append(
        styled_table(
            [["Loại", "URL"]] + attention,
            [35 * mm, doc.width - 35 * mm],
            font_name,
            bold_name,
            wrap=True,
        )
    )

    story.append(PageBreak())
    story.append(Paragraph("ROADMAP 30/60/90 NGÀY", styles["SEO_H1"]))
    roadmap_rows = [["Ngày", "Nhóm", "Ưu tiên", "Hành động", "URL", "KPI/tiêu chí"]]
    roadmap_rows.extend(
        [
            row.get("phase_days", ""),
            row.get("workstream", ""),
            row.get("priority", ""),
            row.get("action", ""),
            row.get("url", ""),
            row.get("success_metric", ""),
        ]
        for row in payload["roadmap"][:45]
    )
    story.append(
        styled_table(
            roadmap_rows,
            [15 * mm, 30 * mm, 16 * mm, 64 * mm, 67 * mm, 59 * mm],
            font_name,
            bold_name,
            wrap=True,
        )
    )
    story.append(Spacer(1, 5 * mm))
    story.append(
        Paragraph(
            "Gợi ý vận hành: xử lý P0 trong 30 ngày, xác minh bằng audit mới; "
            "mở rộng content/authority trong 60 ngày; đánh giá tăng trưởng GSC, "
            "rank và doanh thu ở mốc 90 ngày.",
            body,
        )
    )

    doc.build(story)
    return path


def clean_pdf_text(value):
    text = str(value or "")
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def format_number(value):
    try:
        return f"{float(value):,.0f}"
    except (TypeError, ValueError):
        return str(value)


def styled_table(rows, widths, font_name, bold_name, wrap=False):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, Table, TableStyle

    if not rows:
        rows = [["Không có dữ liệu"]]
        widths = [sum(widths)]
    normal = ParagraphStyle(
        "Cell",
        fontName=font_name,
        fontSize=7.2,
        leading=9.2,
        textColor=colors.HexColor("#" + NAVY),
    )
    header = ParagraphStyle(
        "CellHeader",
        fontName=bold_name,
        fontSize=7.4,
        leading=9.4,
        textColor=colors.white,
    )
    cooked = []
    for row_index, row in enumerate(rows):
        cooked.append(
            [
                Paragraph(clean_pdf_text(value), header if row_index == 0 else normal)
                if wrap or row_index == 0
                else value
                for value in row
            ]
        )
    table = Table(cooked, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BLUE)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), bold_name),
                ("FONTNAME", (0, 1), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 7.2),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D6DEE7")),
                ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#" + LIGHT)]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table
