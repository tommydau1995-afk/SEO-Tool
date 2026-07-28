import tempfile
import unittest
import json
from datetime import date, timedelta
from pathlib import Path

from seo_v4_core import FixItem, PageResult, SerpResearchResult, SerpSource
from seo_v43_core import (
    AdvancedServerLogAnalyzer,
    BacklinkAnalyzer,
    DirectGoogleResearcher,
    ProjectStoreV43,
    analyze_content_decay,
    analyze_content_gap,
    build_link_map,
    build_roadmap,
    competitor_share_of_voice,
)
from seo_v43_reports import export_excel, export_pdf, report_payload


def page(url, depth=0, inlinks=0, internal_links=1, dead_end=False):
    return PageResult(
        score=80,
        status=200,
        url=url,
        final_url=url,
        title="Demo title for SEO",
        title_len=18,
        meta="Demo meta",
        meta_len=9,
        h1="Demo",
        h1_count=1,
        h2_count=2,
        words=500,
        images=1,
        missing_alt=0,
        internal_links=internal_links,
        external_links=0,
        canonical=url,
        indexable="Yes",
        response_ms=200,
        schema_count=1,
        og_title="Demo",
        issues=[],
        depth=depth,
        inlinks=inlinks,
        dead_end=dead_end,
    )


class DirectGoogleTests(unittest.TestCase):
    def test_parses_organic_results_questions_and_ai_overview(self):
        html = """
        <html><body>
          <div><a href="https://one.example/a"><h3>Kết quả một</h3></a>
            <span>Đây là đoạn mô tả đủ dài để trình phân tích nhận dạng snippet.</span>
          </div>
          <div><a href="/url?q=https%3A%2F%2Ftwo.example%2Fb"><h3>Kết quả hai</h3></a></div>
          <div data-q="SEO là gì?">SEO là gì?</div>
          <section data-attrid="AIOverview">Đây là phần tổng quan AI có nội dung dài,
          giải thích chi tiết và đưa ra các điểm chính để người đọc có thể hiểu chủ đề.</section>
        </body></html>
        """
        result = DirectGoogleResearcher.parse_search_html(html)
        self.assertEqual(len(result["results"]), 2)
        self.assertEqual(result["results"][0]["position"], 1)
        self.assertEqual(result["results"][1]["link"], "https://two.example/b")
        self.assertEqual(result["questions"], ["SEO là gì?"])
        self.assertIn("tổng quan AI", result["ai_overview"])


class LinkAndLogTests(unittest.TestCase):
    def test_internal_link_map_and_sitemap_coverage(self):
        home = page("https://example.com/", depth=0)
        article = page("https://example.com/article", depth=1, inlinks=1, dead_end=True)
        graph = {home.url: [article.url], article.url: []}
        result = build_link_map(
            [home, article],
            graph,
            [home.url, article.url, "https://example.com/orphan"],
        )
        self.assertEqual(len(result["edges"]), 1)
        self.assertEqual(result["sitemap_coverage"], 66.7)
        self.assertEqual(result["orphans"], ["https://example.com/orphan"])
        self.assertEqual(result["dead_ends"], [article.url])

    def test_advanced_log_uses_paths_for_known_url_matching(self):
        rows = [
            '1.1.1.1 - - [28/Jul/2026:10:00:00 +0700] "GET /a HTTP/1.1" 200 100 "-" "Googlebot/2.1"',
            '1.1.1.1 - - [28/Jul/2026:10:01:00 +0700] "GET /a HTTP/1.1" 200 100 "-" "Googlebot/2.1"',
            '1.1.1.2 - - [28/Jul/2026:10:02:00 +0700] "GET /search?q=x HTTP/1.1" 404 10 "-" "Mozilla/5.0 Android Googlebot/2.1 Mobile"',
        ]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "access.log"
            path.write_text("\n".join(rows), encoding="utf-8")
            report = AdvancedServerLogAnalyzer().analyze(
                path,
                ["https://example.com/a", "https://example.com/never-crawled"],
            )
        self.assertEqual(report["googlebot_hits"], 3)
        self.assertEqual(report["waste_hits"], 1)
        self.assertEqual(report["crawl_efficiency"], 66.7)
        low = {row["url"]: row["hits"] for row in report["low_crawl_urls"]}
        self.assertNotIn("https://example.com/a", low)
        self.assertEqual(low["https://example.com/never-crawled"], 0)


class GrowthTests(unittest.TestCase):
    def test_backlink_import_and_risk_summary(self):
        csv_text = (
            "Referring Page,Target URL,Anchor Text,DR,Link Type,Organic Traffic\n"
            "https://good.example/a,https://site.example/,brand,55,dofollow,1000\n"
            "https://casino-spam.xyz/x,https://site.example/,casino bonus,2,nofollow,0\n"
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "links.csv"
            path.write_text(csv_text, encoding="utf-8")
            rows = BacklinkAnalyzer.import_csv(path)
        summary = BacklinkAnalyzer.summarize(rows)
        self.assertEqual(summary["backlinks"], 2)
        self.assertEqual(summary["referring_domains"], 2)
        self.assertEqual(summary["dofollow_percent"], 50.0)
        self.assertEqual(summary["high_risk"], 1)

    def test_content_decay_compares_consecutive_30_day_windows(self):
        today = date.today()
        rows = [
            {
                "date": str(today - timedelta(days=5)),
                "page": "https://example.com/a",
                "clicks": "20",
                "impressions": "200",
            },
            {
                "date": str(today - timedelta(days=35)),
                "page": "https://example.com/a",
                "clicks": "100",
                "impressions": "1000",
            },
        ]
        result = analyze_content_decay(rows, reference_date=today)
        self.assertEqual(result[0]["click_change_percent"], -80.0)
        self.assertEqual(result[0]["status"], "Decay nặng")

    def test_content_gap_and_share_of_voice(self):
        sources = [
            SerpSource(
                1,
                "A",
                "https://a.example",
                headings=["H2: Chi phí SEO", "H2: Quy trình SEO"],
            ),
            SerpSource(
                2,
                "B",
                "https://b.example",
                headings=["H2: Chi phí SEO", "H2: Công cụ SEO"],
            ),
        ]
        serp = SerpResearchResult(
            "seo",
            "now",
            sources,
            "",
            [],
            [],
            "",
        )
        gap = analyze_content_gap(
            "seo", "https://site.example/seo", serp, ["H1: Dịch vụ SEO"]
        )
        self.assertEqual(gap["gaps"][0]["topic"], "Chi phí SEO")
        self.assertEqual(gap["gaps"][0]["competitor_count"], 2)
        sov = competitor_share_of_voice(
            [
                {"keyword": "seo", "domain": "a.example", "position": 1},
                {"keyword": "seo", "domain": "b.example", "position": 10},
            ]
        )
        self.assertGreater(sov[0]["share_of_voice"], sov[1]["share_of_voice"])

    def test_roadmap_assigns_30_60_90_day_phases(self):
        fixes = [
            FixItem("P0", "HTTP 500", "https://example.com/a", "Fix server"),
            FixItem("P1", "Thiếu Meta", "https://example.com/b", "Write meta"),
            FixItem("P2", "Thiếu ALT", "https://example.com/c", "Write alt"),
        ]
        result = build_roadmap(fixes, gsc_rows=[{"clicks": 1}])
        self.assertEqual([row["phase_days"] for row in result], [30, 60, 90, 90])

    def test_v43_snapshot_is_stamped_with_new_version(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ProjectStoreV43(directory)
            project = store.upsert_project("Demo", "https://example.com")
            snapshot = store.save_snapshot(
                project, [page("https://example.com/")], {"robots.txt": {}}
            )
            payload = json.loads(snapshot.read_text(encoding="utf-8"))
        self.assertEqual(payload["version"], "4.3.0")


class ReportTests(unittest.TestCase):
    def test_excel_and_pdf_reports_are_created(self):
        payload = report_payload(
            project={"name": "Demo", "website": "https://example.com"},
            pages=[page("https://example.com/")],
            fixes=[
                FixItem(
                    "P0",
                    "Thiếu Title",
                    "https://example.com/",
                    "Viết title",
                )
            ],
            link_map={
                "edges": [],
                "sitemap_urls": 1,
                "sitemap_coverage": 100,
                "orphans": [],
                "dead_ends": [],
                "deep_pages": [],
            },
            roadmap=[
                {
                    "phase_days": 30,
                    "workstream": "Technical SEO",
                    "priority": "P0",
                    "action": "Sửa title",
                    "url": "https://example.com/",
                    "success_metric": "Title hợp lệ",
                    "status": "Cần làm",
                }
            ],
        )
        with tempfile.TemporaryDirectory() as directory:
            xlsx = export_excel(Path(directory) / "report.xlsx", payload)
            pdf = export_pdf(Path(directory) / "report.pdf", payload)
            self.assertGreater(xlsx.stat().st_size, 5000)
            self.assertGreater(pdf.stat().st_size, 3000)
            from openpyxl import load_workbook

            workbook = load_workbook(xlsx, data_only=False)
            self.assertIn("Dashboard", workbook.sheetnames)
            self.assertIn("Roadmap", workbook.sheetnames)
            self.assertTrue(str(workbook["Dashboard"]["B18"].value).startswith("="))


if __name__ == "__main__":
    unittest.main()
