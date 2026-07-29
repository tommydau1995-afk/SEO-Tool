import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from generate_user_guide_v47 import generate
from seo_v4_core import PageResult
from seo_v43_core import FriendlyError
from seo_v47_keywords import (
    GoogleSuggestClient,
    ProjectStoreV47,
    aggregate_gsc_queries,
    discovery_summary,
    discovery_to_keyword_rows,
    export_discovery_csv,
    export_discovery_workbook,
    merge_keyword_candidates,
    onboarding_statuses_v47,
    opportunity_score,
    split_seed_keywords,
    website_keyword_candidates,
)


def page(url="https://example.com/noi-that", title="Nội thất cao cấp | Brand"):
    return PageResult(
        score=90,
        status=200,
        url=url,
        final_url=url,
        title=title,
        title_len=len(title),
        meta="Thiết kế nội thất cao cấp.",
        meta_len=28,
        h1="Thiết kế nội thất cao cấp",
        h1_count=1,
        h2_count=3,
        words=900,
        images=4,
        missing_alt=0,
        internal_links=8,
        external_links=1,
        canonical=url,
        indexable="Yes",
        response_ms=300,
        schema_count=1,
        og_title=title,
    )


class FakeResponse:
    def __init__(self, payload=None, status=200):
        self.payload = payload
        self.status_code = status

    def json(self):
        return self.payload

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")


class FakeSession:
    def __init__(self, response):
        self.response = response
        self.headers = {}
        self.calls = []

    def get(self, url, params=None, timeout=None):
        self.calls.append((url, params, timeout))
        return self.response


class DiscoveryInputTests(unittest.TestCase):
    def test_seed_parser_deduplicates_accents_case_and_delimiters(self):
        self.assertEqual(
            split_seed_keywords(
                "SEO, nội thất cao cấp\nseo; Thiết kế nội thất, nội thất cao cấp"
            ),
            ["SEO", "nội thất cao cấp", "Thiết kế nội thất"],
        )

    def test_google_suggest_payload_is_normalized(self):
        payload = [
            "noi that",
            ["nội thất cao cấp", " nội thất cao cấp ", "https://bad.example"],
        ]
        self.assertEqual(
            GoogleSuggestClient.parse_payload(payload), ["nội thất cao cấp"]
        )

    def test_google_suggest_uses_locale(self):
        session = FakeSession(FakeResponse(["seed", ["gợi ý một", "gợi ý hai"]]))
        rows = GoogleSuggestClient(session=session).suggest(
            "nội thất", country="vn", language="vi"
        )
        self.assertEqual(rows, ["gợi ý một", "gợi ý hai"])
        params = session.calls[0][1]
        self.assertEqual(params["gl"], "vn")
        self.assertEqual(params["hl"], "vi")

    def test_google_suggest_429_has_actionable_error(self):
        session = FakeSession(FakeResponse([], status=429))
        with self.assertRaises(FriendlyError) as context:
            GoogleSuggestClient(session=session).suggest("seo")
        self.assertIn("429", context.exception.message)
        self.assertIn("Nhanh", context.exception.hint)

    def test_expansion_levels_are_bounded(self):
        fast = GoogleSuggestClient.expansion_queries("seo", "vi", "Nhanh")
        deep = GoogleSuggestClient.expansion_queries("seo", "vi", "Sâu")
        self.assertLess(len(fast), len(deep))
        self.assertLessEqual(len(deep), 46)


class DataFusionTests(unittest.TestCase):
    def test_gsc_rows_are_aggregated_without_becoming_volume(self):
        result = aggregate_gsc_queries(
            [
                {
                    "query": "nội thất cao cấp",
                    "page": "https://example.com/noi-that",
                    "clicks": 2,
                    "impressions": 100,
                    "position": 12,
                },
                {
                    "query": "nội thất cao cấp",
                    "page": "https://example.com/noi-that",
                    "clicks": 1,
                    "impressions": 50,
                    "position": 8,
                },
            ]
        )[0]
        self.assertEqual(result["gsc_impressions"], 150)
        self.assertEqual(result["gsc_clicks"], 3)
        self.assertAlmostEqual(result["gsc_position"], 10.7, places=1)
        self.assertNotIn("volume", result)

    def test_website_candidates_include_page_topics(self):
        rows = website_keyword_candidates([page()])
        keywords = {row["keyword"] for row in rows}
        self.assertIn("Thiết kế nội thất cao cấp", keywords)
        self.assertIn("Nội thất cao cấp", keywords)
        self.assertTrue(all(row["landing_page"] for row in rows))

    def test_merge_deduplicates_and_preserves_sources(self):
        result = merge_keyword_candidates(
            suggest_rows=[
                {
                    "keyword": "Nội thất cao cấp",
                    "source": "Google Suggest",
                    "seed": "nội thất",
                }
            ],
            gsc_rows=[
                {
                    "keyword": "nội thất cao cấp",
                    "source": "Google Search Console",
                    "gsc_impressions": 300,
                    "gsc_clicks": 5,
                    "gsc_position": 11,
                    "landing_page": "https://example.com/noi-that",
                }
            ],
            pages=[page()],
        )
        self.assertEqual(len(result), 1)
        self.assertIn("Google Suggest", result[0]["source"])
        self.assertIn("Google Search Console", result[0]["source"])
        self.assertEqual(result[0]["gsc_impressions"], 300)
        self.assertGreaterEqual(result[0]["opportunity"], 70)

    def test_opportunity_is_transparent_not_keyword_difficulty(self):
        score, reason = opportunity_score(
            {
                "keyword": "giá nội thất cao cấp",
                "source": "Google Suggest + Google Search Console",
                "gsc_impressions": 1000,
                "gsc_clicks": 3,
                "gsc_position": 13,
                "intent": "Transactional",
                "landing_page": "https://example.com/noi-that",
            }
        )
        self.assertGreaterEqual(score, 70)
        self.assertIn("impression GSC", reason)

    def test_conversion_never_invents_volume_or_kd(self):
        rows = discovery_to_keyword_rows(
            [
                {
                    "keyword": "giá nội thất cao cấp",
                    "source": "Google Suggest",
                    "intent": "Transactional",
                    "cluster": "Nội Thất",
                    "opportunity": 82,
                    "landing_page": "",
                }
            ]
        )
        self.assertEqual(rows[0]["volume"], 0)
        self.assertEqual(rows[0]["difficulty"], 0)
        self.assertEqual(rows[0]["priority"], "P0")
        self.assertIn("Chưa có Volume/KD", rows[0]["status"])


class PersistenceAndExportTests(unittest.TestCase):
    def test_v47_store_uses_distinct_recovery_and_settings(self):
        with tempfile.TemporaryDirectory() as temp:
            store = ProjectStoreV47(temp)
            store.save_app_settings({"mode": "Easy"})
            store.begin_recovery("Keyword Discovery", {"seeds": ["seo"]})
            self.assertTrue((Path(temp) / "settings-v47.json").exists())
            self.assertTrue((Path(temp) / "recovery-v47.json").exists())
            self.assertEqual(store.load_recovery()["payload"]["seeds"], ["seo"])

    def test_csv_and_excel_exports_keep_data_labels(self):
        rows = [
            {
                "keyword": "nội thất cao cấp",
                "opportunity": 75,
                "opportunity_reason": "GSC",
                "intent": "Commercial",
                "cluster": "Nội Thất",
                "source": "Google Search Console",
                "gsc_impressions": 500,
                "metrics_status": "Chưa có Volume/KD thật",
            }
        ]
        with tempfile.TemporaryDirectory() as temp:
            csv_path = export_discovery_csv(Path(temp) / "keywords.csv", rows)
            with csv_path.open(encoding="utf-8-sig", newline="") as handle:
                exported = list(csv.DictReader(handle))
            self.assertEqual(exported[0]["opportunity"], "75")
            xlsx_path = export_discovery_workbook(
                Path(temp) / "keywords.xlsx", rows
            )
            workbook = load_workbook(xlsx_path, read_only=True)
            self.assertIn("Tổng quan", workbook.sheetnames)
            self.assertIn("Keywords", workbook.sheetnames)
            self.assertIn("V4.7", workbook["Tổng quan"]["A1"].value)
            workbook.close()

    def test_onboarding_replaces_semrush_with_discovery(self):
        rows = onboarding_statuses_v47(
            {"website": "https://example.com"},
            [page()],
            [],
            [{"keyword": "seo"}],
            [{"query": "seo"}],
            {"lab": {"score": 90}},
            True,
        )
        names = [row["name"] for row in rows]
        self.assertIn("Tìm bộ từ khóa", names)
        self.assertNotIn("Import Semrush", names)
        self.assertEqual(
            next(row["status"] for row in rows if row["name"] == "Tìm bộ từ khóa"),
            "Đã kết nối",
        )

    def test_summary_counts_high_opportunity_and_mapping(self):
        summary = discovery_summary(
            [
                {
                    "keyword": "a",
                    "cluster": "A",
                    "opportunity": 80,
                    "source": "Google Search Console",
                    "landing_page": "https://example.com/a",
                },
                {
                    "keyword": "b",
                    "cluster": "B",
                    "opportunity": 50,
                    "source": "Google Suggest",
                    "landing_page": "",
                },
            ]
        )
        self.assertEqual(summary["keywords"], 2)
        self.assertEqual(summary["high_opportunity"], 1)
        self.assertEqual(summary["mapped"], 1)

    def test_v47_guide_documents_keyword_discovery_and_metric_limits(self):
        with tempfile.TemporaryDirectory() as temp:
            path = generate(Path(temp) / "guide-v47.pdf")
            reader = PdfReader(path)
            self.assertGreaterEqual(len(reader.pages), 28)
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
            for token in (
                "Tìm bộ từ khóa không cần Semrush",
                "Google Suggest",
                "Google Search Console",
                "Cơ hội 0-100",
                "Không phải",
                "Volume/KD",
                "Content Brief",
            ):
                self.assertIn(token, text)


if __name__ == "__main__":
    unittest.main()
