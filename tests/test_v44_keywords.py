import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from generate_user_guide_v44 import generate
from seo_v4_core import PageResult
from seo_v44_keywords import (
    ProjectStoreV44,
    SemrushKeywordImporter,
    classify_intent,
    cluster_label,
    cluster_summary,
    export_keyword_workbook,
    keyword_kpis,
    parse_metric,
    parse_decimal_metric,
    rebuild_keyword_intelligence,
)


def sample_page():
    return PageResult(
        score=85,
        status=200,
        url="https://example.com/",
        final_url="https://example.com/",
        title="Demo",
        title_len=4,
        meta="Demo",
        meta_len=4,
        h1="Demo",
        h1_count=1,
        h2_count=2,
        words=500,
        images=1,
        missing_alt=0,
        internal_links=2,
        external_links=0,
        canonical="https://example.com/",
        indexable="Yes",
        response_ms=250,
        schema_count=1,
        og_title="Demo",
        issues=[],
    )


class MetricAndIntentTests(unittest.TestCase):
    def test_metric_parser_handles_semrush_number_formats(self):
        self.assertEqual(parse_metric("1,200"), 1200)
        self.assertEqual(parse_metric("1.200"), 1200)
        self.assertEqual(parse_metric("1,2K"), 1200)
        self.assertEqual(parse_metric("1.2K"), 1200)
        self.assertEqual(parse_metric("62%"), 62)
        self.assertEqual(parse_decimal_metric("1.200"), 1.2)
        self.assertEqual(parse_decimal_metric("0,25"), 0.25)

    def test_search_intent_classifier(self):
        self.assertEqual(classify_intent("seo là gì"), "Informational")
        self.assertEqual(classify_intent("top công ty seo"), "Commercial")
        self.assertEqual(
            classify_intent("báo giá dịch vụ seo"), "Transactional"
        )
        self.assertEqual(classify_intent("semrush login"), "Navigational")

    def test_cluster_detail_levels(self):
        keyword = "báo giá dịch vụ seo tổng thể website"
        self.assertEqual(cluster_label(keyword, "Broad"), "Seo")
        self.assertEqual(cluster_label(keyword, "Balanced"), "Seo Tổng")
        self.assertEqual(cluster_label(keyword, "Tight"), "Seo Tổng Thể")


class SemrushImporterTests(unittest.TestCase):
    def test_csv_import_detects_columns_and_removes_duplicates(self):
        value = (
            "Keyword;Search Volume;KD %;Intent;CPC (USD);Position;Ranking URL\n"
            "seo là gì;1.200;45;I;0,25;3;https://example.com/seo\n"
            "SEO LÀ GÌ;900;42;Informational;0,20;5;https://example.com/seo\n"
            "báo giá dịch vụ seo;1,2K;61;T;2,50;;\n"
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "semrush.csv"
            path.write_text(value, encoding="utf-8-sig")
            rows, summary = SemrushKeywordImporter().read(path)
        self.assertEqual(summary["rows_read"], 3)
        self.assertEqual(summary["keywords_imported"], 2)
        self.assertEqual(summary["duplicates_removed"], 1)
        self.assertEqual(summary["detected"]["volume"], "Search Volume")
        self.assertEqual(rows[0]["volume"], 1200)
        self.assertEqual(rows[0]["difficulty"], 45)
        self.assertEqual(rows[1]["intent"], "Transactional")
        self.assertEqual(rows[1]["cluster"], "Seo")

    def test_xlsx_import_supports_vietnamese_aliases(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "keyword.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Từ khóa", "Lượng tìm kiếm", "Độ khó", "Ý định"])
            sheet.append(["cách làm seo", 500, 35, "Thông tin"])
            workbook.save(path)
            rows, summary = SemrushKeywordImporter().read(path)
        self.assertEqual(summary["keywords_imported"], 1)
        self.assertEqual(rows[0]["keyword"], "cách làm seo")
        self.assertEqual(rows[0]["volume"], 500)
        self.assertEqual(rows[0]["intent"], "Informational")

    def test_rebuild_and_cluster_summary(self):
        rows = rebuild_keyword_intelligence(
            [
                {
                    "keyword": "seo là gì",
                    "volume": 1000,
                    "difficulty": 40,
                },
                {
                    "keyword": "dịch vụ seo",
                    "volume": 800,
                    "difficulty": 55,
                },
                {
                    "keyword": "seo website",
                    "volume": 400,
                    "difficulty": 35,
                },
            ],
            "Balanced",
        )
        summary = cluster_summary(rows)
        seo = next(row for row in summary if row["cluster"] == "Seo")
        self.assertEqual(seo["keywords"], 2)
        self.assertEqual(seo["total_volume"], 1800)
        kpis = keyword_kpis(rows)
        self.assertEqual(kpis["keywords"], 3)
        self.assertEqual(kpis["total_volume"], 2200)
        self.assertEqual(kpis["clusters"], 2)


class V44ArtifactTests(unittest.TestCase):
    def test_keyword_excel_has_dashboard_and_cluster_sheets(self):
        rows = rebuild_keyword_intelligence(
            [
                {
                    "keyword": "dịch vụ seo",
                    "volume": 1200,
                    "difficulty": 50,
                },
                {
                    "keyword": "seo là gì",
                    "volume": 900,
                    "difficulty": 35,
                },
            ]
        )
        with tempfile.TemporaryDirectory() as directory:
            path = export_keyword_workbook(
                Path(directory) / "keyword-intelligence.xlsx", rows
            )
            workbook = load_workbook(path, data_only=False)
        self.assertEqual(
            workbook.sheetnames, ["Tổng quan", "Keywords", "Clusters"]
        )
        self.assertEqual(workbook["Tổng quan"]["B5"].value, 2100)

    def test_beginner_guide_is_a_multi_page_pdf(self):
        with tempfile.TemporaryDirectory() as directory:
            path = generate(Path(directory) / "guide.pdf")
            content = path.read_bytes()
        self.assertTrue(content.startswith(b"%PDF"))
        self.assertGreater(len(content), 30000)
        self.assertGreaterEqual(content.count(b"/Type /Page"), 10)

    def test_v44_snapshot_version(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ProjectStoreV44(directory)
            project = store.upsert_project("Demo", "https://example.com")
            snapshot = store.save_snapshot(
                project, [sample_page()], {"robots.txt": {"status": 200}}
            )
            payload = json.loads(snapshot.read_text(encoding="utf-8"))
        self.assertEqual(payload["version"], "4.4.0")


if __name__ == "__main__":
    unittest.main()
