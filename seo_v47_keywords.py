"""Built-in keyword discovery for SEO AI Studio V4.7.

The engine deliberately separates discovery signals from paid keyword metrics:

* Google Suggest is used only to discover real query phrasing.
* Search Console contributes first-party impressions, clicks and positions.
* Crawled pages contribute existing topics and candidate landing pages.
* Volume and KD remain empty unless the user imports a real metrics source.

This keeps the default workflow useful without Semrush while avoiding invented
search-volume or keyword-difficulty values.
"""

from __future__ import annotations

import math
import csv
import json
import re
import string
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable
from urllib.parse import unquote, urlparse

import requests

from seo_v4_core import PageResult, clean_text, normalize_url
from seo_v43_core import FriendlyError, parse_number
from seo_v44_keywords import (
    ascii_key,
    classify_intent,
    cluster_label,
    deduplicate_keywords,
    normalize_keyword_row,
)
from seo_v46_core import ProjectStoreV46


VERSION = "4.7.0"
GUIDE_FILE = "Huong_Dan_Chi_Tiet_SEO_AI_Studio_V47.pdf"
UPDATE_MANIFEST_URL = (
    "https://raw.githubusercontent.com/tommydau1995-afk/SEO-Tool/"
    "main/update_manifest_v47.json"
)

DISCOVERY_LEVELS = {
    "Nhanh": {"alphabet": "", "questions": 3, "max_requests_per_seed": 4},
    "Cân bằng": {"alphabet": "abcdefgilmnoprstuv", "questions": 6, "max_requests_per_seed": 18},
    "Sâu": {"alphabet": string.ascii_lowercase + "0123456789", "questions": 9, "max_requests_per_seed": 46},
}

QUESTION_PREFIXES = {
    "vi": (
        "cách",
        "là gì",
        "tại sao",
        "ở đâu",
        "bao nhiêu",
        "có nên",
        "kinh nghiệm",
        "so sánh",
        "review",
    ),
    "en": (
        "how to",
        "what is",
        "why",
        "where",
        "how much",
        "best",
        "review",
        "compare",
        "near me",
    ),
}

GENERIC_PAGE_TOKENS = {
    "trang chu",
    "home",
    "homepage",
    "lien he",
    "contact",
    "gioi thieu",
    "about us",
    "dieu khoan",
    "privacy policy",
}


def split_seed_keywords(value: str, limit: int = 10) -> list[str]:
    """Return unique seed keywords from lines, commas or semicolons."""

    result: list[str] = []
    seen: set[str] = set()
    for raw in re.split(r"[\n,;]+", value or ""):
        keyword = clean_text(raw)
        key = ascii_key(keyword)
        if not key or key in seen:
            continue
        seen.add(key)
        result.append(keyword)
        if len(result) >= max(1, int(limit)):
            break
    return result


def normalize_candidate_keyword(value: str) -> str:
    value = clean_text(value).strip(" -–—|,.;:")
    value = re.sub(r"\s+", " ", value)
    if len(value) < 2 or len(value) > 120:
        return ""
    if value.startswith(("http://", "https://")):
        return ""
    return value


def _page_phrases(page: PageResult | dict) -> list[str]:
    def get(field: str):
        return (
            getattr(page, field, "")
            if not isinstance(page, dict)
            else page.get(field, "")
        )

    phrases: list[str] = []
    for value in (get("h1"), get("title")):
        value = clean_text(value)
        if not value:
            continue
        # Brand names often follow a pipe/dash in page titles.
        value = re.split(r"\s+[|–—]\s+|\s+-\s+", value, maxsplit=1)[0]
        value = normalize_candidate_keyword(value)
        key = ascii_key(value)
        if value and key not in GENERIC_PAGE_TOKENS and 2 <= len(value.split()) <= 12:
            phrases.append(value)

    url = normalize_url(get("url"))
    if url:
        slug = unquote(urlparse(url).path.rstrip("/").split("/")[-1])
        slug = normalize_candidate_keyword(re.sub(r"[-_]+", " ", slug))
        key = ascii_key(slug)
        if slug and key not in GENERIC_PAGE_TOKENS and 2 <= len(slug.split()) <= 10:
            phrases.append(slug)
    return list(dict.fromkeys(phrases))


def website_keyword_candidates(pages: Iterable[PageResult | dict]) -> list[dict]:
    rows = []
    for page in pages or ():
        url = normalize_url(
            getattr(page, "url", "") if not isinstance(page, dict) else page.get("url")
        )
        for keyword in _page_phrases(page):
            rows.append(
                {
                    "keyword": keyword,
                    "source": "Website",
                    "existing_url": url,
                    "landing_page": url,
                }
            )
    return rows


def aggregate_gsc_queries(rows: Iterable[dict]) -> list[dict]:
    """Aggregate date/page/query Search Console rows without treating impressions as volume."""

    grouped: dict[str, dict] = {}
    for row in rows or ():
        keyword = normalize_candidate_keyword(row.get("query", ""))
        key = ascii_key(keyword)
        if not key:
            continue
        clicks = max(0.0, parse_number(row.get("clicks", 0)))
        impressions = max(0.0, parse_number(row.get("impressions", 0)))
        position = max(0.0, parse_number(row.get("position", 0)))
        page = normalize_url(row.get("page", ""))
        item = grouped.setdefault(
            key,
            {
                "keyword": keyword,
                "source": "Google Search Console",
                "gsc_clicks": 0.0,
                "gsc_impressions": 0.0,
                "_position_weight": 0.0,
                "_page_impressions": defaultdict(float),
            },
        )
        item["gsc_clicks"] += clicks
        item["gsc_impressions"] += impressions
        item["_position_weight"] += position * max(1.0, impressions)
        if page:
            item["_page_impressions"][page] += max(1.0, impressions)

    result = []
    for item in grouped.values():
        denominator = max(1.0, item["gsc_impressions"])
        page_weights = item.pop("_page_impressions")
        item["gsc_position"] = round(item.pop("_position_weight") / denominator, 1)
        item["gsc_clicks"] = round(item["gsc_clicks"], 1)
        item["gsc_impressions"] = round(item["gsc_impressions"], 1)
        item["existing_url"] = (
            max(page_weights, key=page_weights.get) if page_weights else ""
        )
        item["landing_page"] = item["existing_url"]
        result.append(item)
    return result


class GoogleSuggestClient:
    """Best-effort Google autocomplete client; no API key is required."""

    ENDPOINT = "https://suggestqueries.google.com/complete/search"

    def __init__(self, timeout: int = 12, session=None):
        self.timeout = max(3, int(timeout))
        self.http = session or requests.Session()
        self.http.headers.update(
            {
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 Chrome/126 Safari/537.36 SEOAIStudio/4.7"
                ),
                "Accept-Language": "vi,en-US;q=0.8,en;q=0.7",
            }
        )

    @staticmethod
    def parse_payload(payload) -> list[str]:
        if not isinstance(payload, list) or len(payload) < 2:
            return []
        suggestions = payload[1]
        if not isinstance(suggestions, list):
            return []
        result = []
        seen = set()
        for raw in suggestions:
            keyword = normalize_candidate_keyword(raw)
            key = ascii_key(keyword)
            if keyword and key not in seen:
                seen.add(key)
                result.append(keyword)
        return result

    def suggest(self, query: str, country: str = "vn", language: str = "vi") -> list[str]:
        query = normalize_candidate_keyword(query)
        if not query:
            return []
        try:
            response = self.http.get(
                self.ENDPOINT,
                params={
                    "client": "firefox",
                    "q": query,
                    "gl": clean_text(country).lower() or "vn",
                    "hl": clean_text(language).lower() or "vi",
                },
                timeout=self.timeout,
            )
            if response.status_code == 429:
                raise FriendlyError(
                    "Google Suggest giới hạn tần suất",
                    "Google tạm thời trả HTTP 429 vì có quá nhiều yêu cầu.",
                    "Chọn mức Nhanh, giảm số từ khóa gốc và thử lại sau vài phút.",
                )
            response.raise_for_status()
            return self.parse_payload(response.json())
        except FriendlyError:
            raise
        except requests.Timeout as exc:
            raise FriendlyError(
                "Google Suggest quá thời gian",
                "Không nhận được gợi ý trong thời gian cho phép.",
                "Kiểm tra Internet, chọn mức Nhanh rồi nhấn Thử lại.",
            ) from exc
        except (requests.RequestException, ValueError) as exc:
            raise FriendlyError(
                "Không lấy được Google Suggest",
                clean_text(exc) or "Google Suggest trả dữ liệu không hợp lệ.",
                "Kiểm tra Internet hoặc dùng dữ liệu GSC/Website đang có.",
            ) from exc

    @staticmethod
    def expansion_queries(seed: str, language: str, level: str) -> list[str]:
        config = DISCOVERY_LEVELS.get(level, DISCOVERY_LEVELS["Cân bằng"])
        prefixes = QUESTION_PREFIXES.get(
            clean_text(language).lower(), QUESTION_PREFIXES["en"]
        )
        queries = [seed]
        for prefix in prefixes[: config["questions"]]:
            if prefix in ("là gì", "bao nhiêu", "có nên", "near me"):
                queries.append(f"{seed} {prefix}")
            else:
                queries.append(f"{prefix} {seed}")
        queries.extend(f"{seed} {letter}" for letter in config["alphabet"])
        return list(dict.fromkeys(queries))[: config["max_requests_per_seed"]]

    def expand(
        self,
        seeds: Iterable[str],
        country: str = "vn",
        language: str = "vi",
        level: str = "Cân bằng",
        max_keywords: int = 500,
        progress: Callable[[str, int, int], None] | None = None,
        cancelled: Callable[[], bool] | None = None,
    ) -> list[dict]:
        jobs = []
        for seed in seeds:
            jobs.extend(
                (seed, query)
                for query in self.expansion_queries(seed, language, level)
            )
        rows = []
        seen = set()
        for index, (seed, query) in enumerate(jobs, 1):
            if cancelled and cancelled():
                break
            if progress:
                progress(query, index, len(jobs))
            for keyword in self.suggest(query, country, language):
                key = ascii_key(keyword)
                if not key or key in seen:
                    continue
                seen.add(key)
                rows.append(
                    {
                        "keyword": keyword,
                        "seed": seed,
                        "source": "Google Suggest",
                    }
                )
                if len(rows) >= max(1, int(max_keywords)):
                    return rows
            # A small pause reduces accidental rate limiting in deep mode.
            if len(jobs) > 8:
                time.sleep(0.08)
        return rows


def _keyword_tokens(value: str) -> set[str]:
    return {
        token
        for token in re.findall(r"[a-z0-9]+", ascii_key(value))
        if len(token) > 1
    }


def _best_page(keyword: str, pages: Iterable[PageResult | dict]) -> tuple[str, float]:
    keyword_tokens = _keyword_tokens(keyword)
    if not keyword_tokens:
        return "", 0.0
    best_url, best_score = "", 0.0
    for page in pages or ():
        get = (
            (lambda field: page.get(field, ""))
            if isinstance(page, dict)
            else (lambda field: getattr(page, field, ""))
        )
        url = normalize_url(get("url"))
        text = " ".join((clean_text(get("title")), clean_text(get("h1")), unquote(url)))
        page_tokens = _keyword_tokens(text)
        if not page_tokens:
            continue
        overlap = len(keyword_tokens & page_tokens)
        union = len(keyword_tokens | page_tokens)
        score = overlap / max(1, union)
        if score > best_score:
            best_url, best_score = url, score
    return (best_url, round(best_score, 3)) if best_score >= 0.18 else ("", 0.0)


def opportunity_score(row: dict) -> tuple[int, str]:
    """Return a transparent priority score, not an SEO keyword-difficulty metric."""

    score = 20
    reasons = []
    sources = {
        clean_text(value)
        for value in clean_text(row.get("source")).split(" + ")
        if clean_text(value)
    }
    score += min(18, len(sources) * 6)
    if len(sources) >= 2:
        reasons.append("xuất hiện ở nhiều nguồn")

    impressions = max(0.0, parse_number(row.get("gsc_impressions", 0)))
    clicks = max(0.0, parse_number(row.get("gsc_clicks", 0)))
    position = max(0.0, parse_number(row.get("gsc_position", 0)))
    if impressions:
        score += min(22, int(math.log10(impressions + 1) * 7))
        reasons.append(f"{impressions:,.0f} impression GSC")
        if 5 <= position <= 30:
            score += 14
            reasons.append(f"vị trí {position:.1f} có thể cải thiện")
        elif position > 30:
            score += 6
        ctr = clicks / impressions
        if impressions >= 50 and ctr < 0.02:
            score += 8
            reasons.append("CTR còn thấp")

    words = len(clean_text(row.get("keyword")).split())
    if 3 <= words <= 8:
        score += 8
        reasons.append("long-tail rõ ý định")
    intent = clean_text(row.get("intent"))
    if intent in ("Commercial", "Transactional"):
        score += 8
        reasons.append(f"intent {intent}")
    if row.get("landing_page"):
        score += 5
        reasons.append("đã có Landing Page phù hợp")
    else:
        score += 4
        reasons.append("có thể tạo Landing Page mới")

    return min(100, int(round(score))), "; ".join(reasons[:4]) or "Tín hiệu khám phá ban đầu"


def merge_keyword_candidates(
    suggest_rows: Iterable[dict] = (),
    gsc_rows: Iterable[dict] = (),
    website_rows: Iterable[dict] = (),
    pages: Iterable[PageResult | dict] = (),
    cluster_detail: str = "Balanced",
    limit: int = 1000,
) -> list[dict]:
    merged: dict[str, dict] = {}
    order: list[str] = []
    for raw in [*suggest_rows, *gsc_rows, *website_rows]:
        keyword = normalize_candidate_keyword(raw.get("keyword", ""))
        key = ascii_key(keyword)
        if not key:
            continue
        if key not in merged:
            merged[key] = {
                "keyword": keyword,
                "seed": clean_text(raw.get("seed")),
                "source": clean_text(raw.get("source")) or "Khác",
                "gsc_clicks": 0.0,
                "gsc_impressions": 0.0,
                "gsc_position": 0.0,
                "existing_url": "",
                "landing_page": "",
            }
            order.append(key)
        item = merged[key]
        source = clean_text(raw.get("source"))
        current_sources = clean_text(item.get("source")).split(" + ")
        if source and source not in current_sources:
            item["source"] = " + ".join([*current_sources, source])
        for field in ("seed", "existing_url", "landing_page"):
            if not clean_text(item.get(field)) and clean_text(raw.get(field)):
                item[field] = clean_text(raw.get(field))
        if parse_number(raw.get("gsc_impressions", 0)) > parse_number(
            item.get("gsc_impressions", 0)
        ):
            item["gsc_clicks"] = parse_number(raw.get("gsc_clicks", 0))
            item["gsc_impressions"] = parse_number(raw.get("gsc_impressions", 0))
            item["gsc_position"] = parse_number(raw.get("gsc_position", 0))

    result = []
    for key in order:
        item = merged[key]
        if not item["landing_page"]:
            page, similarity = _best_page(item["keyword"], pages)
            if page:
                item["landing_page"] = page
                item["page_similarity"] = similarity
        item["intent"] = classify_intent(item["keyword"])
        item["cluster"] = cluster_label(item["keyword"], cluster_detail)
        item["metrics_status"] = "Chưa có Volume/KD thật"
        item["opportunity"], item["opportunity_reason"] = opportunity_score(item)
        result.append(item)

    result.sort(
        key=lambda row: (
            -int(row.get("opportunity", 0)),
            -parse_number(row.get("gsc_impressions", 0)),
            ascii_key(row.get("keyword")),
        )
    )
    return result[: max(1, int(limit))]


def discovery_to_keyword_rows(rows: Iterable[dict]) -> list[dict]:
    """Convert selected discovery candidates to the established V4 keyword workflow."""

    result = []
    for raw in rows or ():
        score = int(parse_number(raw.get("opportunity", 0)))
        source = "Built-in Discovery: " + clean_text(raw.get("source"))
        row = normalize_keyword_row(
            {
                "keyword": raw.get("keyword"),
                "volume": raw.get("volume", 0),
                "difficulty": raw.get("difficulty", 0),
                "intent": raw.get("intent"),
                "cluster": raw.get("cluster"),
                "url": raw.get("landing_page"),
                "priority": "P0" if score >= 75 else "P1" if score >= 50 else "P2",
                "status": (
                    "Có Landing Page • Chưa có Volume/KD"
                    if raw.get("landing_page")
                    else "Chưa map • Chưa có Volume/KD"
                ),
                "source": source,
            },
            source=source,
        )
        result.append(row)
    return deduplicate_keywords(result)


def discovery_summary(rows: Iterable[dict]) -> dict:
    rows = list(rows or ())
    return {
        "keywords": len(rows),
        "clusters": len({row.get("cluster") for row in rows if row.get("cluster")}),
        "gsc_keywords": sum(
            "Google Search Console" in clean_text(row.get("source")) for row in rows
        ),
        "website_keywords": sum(
            "Website" in clean_text(row.get("source")) for row in rows
        ),
        "high_opportunity": sum(int(row.get("opportunity", 0)) >= 70 for row in rows),
        "mapped": sum(bool(row.get("landing_page")) for row in rows),
    }


DISCOVERY_EXPORT_FIELDS = (
    "keyword",
    "opportunity",
    "opportunity_reason",
    "intent",
    "cluster",
    "source",
    "seed",
    "gsc_clicks",
    "gsc_impressions",
    "gsc_position",
    "landing_page",
    "existing_url",
    "metrics_status",
)


def export_discovery_csv(path, rows: Iterable[dict]) -> Path:
    path = Path(path)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=DISCOVERY_EXPORT_FIELDS, extrasaction="ignore"
        )
        writer.writeheader()
        writer.writerows(rows)
    return path


def export_discovery_workbook(path, rows: Iterable[dict]) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError("Thiếu openpyxl để xuất Excel.") from exc

    rows = list(rows or ())
    summary = discovery_summary(rows)
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Tổng quan"
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:F2")
    dashboard["A1"] = "SEO AI STUDIO V4.7 - KEYWORD DISCOVERY"
    dashboard["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    dashboard["A1"].fill = PatternFill("solid", fgColor="17324D")
    dashboard["A1"].alignment = Alignment(vertical="center")
    metrics = (
        ("Keywords", summary["keywords"]),
        ("Clusters", summary["clusters"]),
        ("Cơ hội ≥70", summary["high_opportunity"]),
        ("Có GSC", summary["gsc_keywords"]),
        ("Có Landing Page", summary["mapped"]),
    )
    for column, (label, value) in enumerate(metrics, 1):
        dashboard.cell(4, column, label)
        dashboard.cell(5, column, value)
        dashboard.cell(4, column).font = Font(bold=True, color="1769C2")
        dashboard.cell(5, column).font = Font(size=16, bold=True, color="17324D")
        dashboard.column_dimensions[dashboard.cell(4, column).column_letter].width = 20
    dashboard["A8"] = "Lưu ý dữ liệu"
    dashboard["A9"] = (
        "Opportunity là điểm ưu tiên nội bộ, không phải Keyword Difficulty. "
        "GSC impressions không phải Search Volume."
    )
    dashboard.merge_cells("A9:F10")
    dashboard["A9"].alignment = Alignment(wrap_text=True, vertical="top")
    dashboard["A12"] = "Tạo lúc"
    dashboard["B12"] = datetime.now().astimezone().isoformat(timespec="seconds")

    sheet = workbook.create_sheet("Keywords")
    sheet.sheet_view.showGridLines = False
    widths = (38, 12, 52, 18, 28, 30, 22, 12, 15, 12, 50, 50, 25)
    labels = (
        "Keyword",
        "Cơ hội",
        "Lý do",
        "Search Intent",
        "Cluster",
        "Nguồn",
        "Seed",
        "GSC Clicks",
        "GSC Impressions",
        "GSC Position",
        "Landing Page",
        "Existing URL",
        "Trạng thái Volume/KD",
    )
    for column, (field, label, width) in enumerate(
        zip(DISCOVERY_EXPORT_FIELDS, labels, widths), 1
    ):
        cell = sheet.cell(1, column, label)
        cell.fill = PatternFill("solid", fgColor="1769C2")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[cell.column_letter].width = width
        for row_index, item in enumerate(rows, 2):
            value = item.get(field, "")
            data_cell = sheet.cell(row_index, column, value)
            data_cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    if rows:
        table = Table(displayName="KeywordDiscoveryTable", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        sheet.add_table(table)
    workbook.calculation.fullCalcOnLoad = True
    workbook.save(path)
    return path


def onboarding_statuses_v47(
    project,
    pages,
    keyword_rows,
    discovery_rows,
    gsc_rows,
    pagespeed,
    wordpress_ok,
):
    def status(condition, error=False):
        return "Có lỗi" if error else "Đã kết nối" if condition else "Chưa thiết lập"

    return [
        {"step": 1, "name": "Tạo dự án", "status": status(bool(project))},
        {
            "step": 2,
            "name": "Nhập website",
            "status": status(bool(normalize_url((project or {}).get("website", "")))),
        },
        {
            "step": 3,
            "name": "Quốc gia & ngôn ngữ",
            "status": "Đã kết nối",
        },
        {"step": 4, "name": "Kiểm tra website", "status": status(bool(pages))},
        {"step": 5, "name": "Google Search Console", "status": status(bool(gsc_rows))},
        {
            "step": 6,
            "name": "PageSpeed/CrUX",
            "status": status(bool(pagespeed), bool(pagespeed.get("error")) if pagespeed else False),
        },
        {"step": 7, "name": "WordPress", "status": status(bool(wordpress_ok))},
        {
            "step": 8,
            "name": "Tìm bộ từ khóa",
            "status": status(bool(discovery_rows or keyword_rows)),
        },
        {"step": 9, "name": "Full SEO Audit", "status": status(bool(pages))},
    ]


class ProjectStoreV47(ProjectStoreV46):
    SETTINGS_FILE = "settings-v47.json"
    RECOVERY_FILE = "recovery-v47.json"

    def save_snapshot(self, project, pages, site_files, keyword_map=None, fixes=None):
        path = super().save_snapshot(
            project, pages, site_files, keyword_map=keyword_map, fixes=fixes
        )
        payload = json.loads(path.read_text(encoding="utf-8"))
        payload["version"] = VERSION
        self._atomic_json(path, payload)
        return path
