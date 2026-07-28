"""Keyword Intelligence for SEO AI Studio V4.4.

Semrush exports vary by locale, account and report type.  The importer accepts
CSV, TSV and XLSX files, maps common column aliases, then enriches every row
with a deterministic Search Intent and editable topic cluster.
"""

from __future__ import annotations

import csv
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from seo_v4_core import clean_text, normalize_url
from seo_v43_core import ProjectStoreV43


VERSION = "4.4.0"

KEYWORD_ALIASES = (
    "keyword",
    "keywords",
    "từ khóa",
    "tu khoa",
    "search term",
    "query",
)
VOLUME_ALIASES = (
    "volume",
    "search volume",
    "monthly volume",
    "avg. monthly searches",
    "lượng tìm kiếm",
    "luong tim kiem",
)
DIFFICULTY_ALIASES = (
    "keyword difficulty",
    "keyword difficulty %",
    "kd",
    "kd %",
    "difficulty",
    "độ khó",
    "do kho",
)
INTENT_ALIASES = ("intent", "search intent", "ý định", "y dinh")
CPC_ALIASES = ("cpc", "cost per click", "cpc (usd)", "giá thầu", "gia thau")
URL_ALIASES = (
    "url",
    "landing page",
    "ranking url",
    "target url",
    "trang đích",
    "trang dich",
)
POSITION_ALIASES = ("position", "pos.", "rank", "ranking", "vị trí", "vi tri")
TREND_ALIASES = ("trend", "trends", "xu hướng", "xu huong")
COMPETITION_ALIASES = (
    "competitive density",
    "competition",
    "com.",
    "mức cạnh tranh",
)
RESULTS_ALIASES = ("results", "number of results", "kết quả", "ket qua")

INTENT_COLORS = {
    "Informational": "#E9F3FF",
    "Commercial": "#FFF2D8",
    "Transactional": "#E8F7EC",
    "Navigational": "#F1EAFE",
}
INTENT_CODES = {
    "i": "Informational",
    "informational": "Informational",
    "information": "Informational",
    "thông tin": "Informational",
    "c": "Commercial",
    "commercial": "Commercial",
    "commercial investigation": "Commercial",
    "thương mại": "Commercial",
    "t": "Transactional",
    "transactional": "Transactional",
    "transaction": "Transactional",
    "giao dịch": "Transactional",
    "n": "Navigational",
    "navigational": "Navigational",
    "navigation": "Navigational",
    "điều hướng": "Navigational",
}

STOPWORDS = {
    "a",
    "an",
    "and",
    "at",
    "best",
    "cach",
    "cac",
    "cho",
    "co",
    "cua",
    "de",
    "duoc",
    "for",
    "gi",
    "how",
    "in",
    "la",
    "lam",
    "mot",
    "nhung",
    "of",
    "o",
    "the",
    "to",
    "top",
    "tren",
    "va",
    "voi",
    "what",
    "where",
    "which",
}
INTENT_MODIFIERS = {
    "bao",
    "buy",
    "cheap",
    "comparison",
    "dich",
    "download",
    "free",
    "gia",
    "guide",
    "huong",
    "mua",
    "nhat",
    "phi",
    "pricing",
    "quy",
    "review",
    "service",
    "tai",
    "tot",
    "trinh",
    "tutorial",
    "vu",
}


class ProjectStoreV44(ProjectStoreV43):
    def save_snapshot(self, project, pages, site_files, keyword_map=None, fixes=None):
        path = super().save_snapshot(project, pages, site_files, keyword_map, fixes)
        import json

        payload = json.loads(path.read_text(encoding="utf-8"))
        payload["version"] = VERSION
        path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return path


def ascii_key(value):
    value = unicodedata.normalize("NFKD", clean_text(value))
    value = value.encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def parse_metric(value):
    text = clean_text(value)
    if not text:
        return 0.0
    text = text.replace("\u00a0", "").replace(" ", "").replace("%", "")
    multiplier = 1
    if text[-1:].lower() == "k":
        multiplier, text = 1_000, text[:-1]
    elif text[-1:].lower() == "m":
        multiplier, text = 1_000_000, text[:-1]
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        text = "".join(parts) if len(parts[-1]) == 3 else text.replace(",", ".")
    elif "." in text:
        parts = text.split(".")
        if len(parts) > 1 and len(parts[-1]) == 3 and all(
            part.isdigit() for part in parts
        ):
            text = "".join(parts)
    text = re.sub(r"[^0-9.\-]", "", text)
    try:
        return round(float(text) * multiplier, 4)
    except ValueError:
        return 0.0


def parse_decimal_metric(value):
    text = clean_text(value)
    if not text:
        return 0.0
    text = text.replace("\u00a0", "").replace(" ", "").replace("%", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    text = re.sub(r"[^0-9.\-]", "", text)
    try:
        return round(float(text), 4)
    except ValueError:
        return 0.0


def _value(row, aliases, default=""):
    normalized = {ascii_key(key): value for key, value in row.items()}
    for alias in aliases:
        key = ascii_key(alias)
        if key in normalized:
            return normalized[key]
    return default


def normalize_intent(value, keyword=""):
    raw = ascii_key(value)
    if raw:
        parts = [part for part in re.split(r"[,/;|]+", raw) if part]
        for part in parts:
            part = part.strip()
            if part in INTENT_CODES:
                return INTENT_CODES[part]
            for token, intent in INTENT_CODES.items():
                if token in part:
                    return intent
    return classify_intent(keyword)


def classify_intent(keyword):
    value = " " + ascii_key(keyword) + " "
    navigation = (
        " login ",
        " dang nhap ",
        " official ",
        " website ",
        " trang chu ",
        " contact ",
        " lien he ",
        " near me ",
    )
    transactional = (
        " mua ",
        " dat ",
        " dang ky ",
        " bao gia ",
        " gia ",
        " thue ",
        " dich vu ",
        " download ",
        " coupon ",
        " khuyen mai ",
        " sale ",
        " pricing ",
        " order ",
        " book ",
    )
    commercial = (
        " review ",
        " danh gia ",
        " so sanh ",
        " tot nhat ",
        " top ",
        " vs ",
        " nen mua ",
        " lua chon ",
        " alternative ",
        " alternatives ",
    )
    informational = (
        " la gi ",
        " tai sao ",
        " cach ",
        " huong dan ",
        " quy trinh ",
        " kinh nghiem ",
        " checklist ",
        " what ",
        " why ",
        " how ",
        " tutorial ",
        " guide ",
    )
    if any(token in value for token in navigation):
        return "Navigational"
    if any(token in value for token in transactional):
        return "Transactional"
    if any(token in value for token in commercial):
        return "Commercial"
    if any(token in value for token in informational):
        return "Informational"
    return "Informational"


def keyword_tokens(keyword):
    display = re.findall(r"[0-9A-Za-zÀ-ỹ]+", clean_text(keyword).lower())
    result = []
    for token in display:
        key = ascii_key(token)
        if not key:
            continue
        if key in STOPWORDS and not (key == "the" and token == "thể"):
            continue
        if key in INTENT_MODIFIERS and not (key == "bao" and token == "bảo"):
            continue
        result.append((key, token))
    return result


def cluster_label(keyword, detail="Balanced"):
    tokens = keyword_tokens(keyword)
    if not tokens:
        return clean_text(keyword).title() or "Chưa phân nhóm"
    size = {"Broad": 1, "Balanced": 2, "Tight": 3}.get(detail, 2)
    chosen = [display for _, display in tokens[:size]]
    return " ".join(chosen).title()


def normalize_keyword_row(row, source="Manual", cluster_detail="Balanced"):
    keyword = clean_text(row.get("keyword") or _value(row, KEYWORD_ALIASES))
    volume = parse_metric(row.get("volume") or _value(row, VOLUME_ALIASES))
    difficulty = parse_decimal_metric(
        row.get("difficulty") or _value(row, DIFFICULTY_ALIASES)
    )
    supplied_intent = row.get("intent") or _value(row, INTENT_ALIASES)
    url = normalize_url(row.get("url") or _value(row, URL_ALIASES))
    position = row.get("rank") or row.get("position") or _value(row, POSITION_ALIASES)
    intent = normalize_intent(supplied_intent, keyword)
    cluster = clean_text(row.get("cluster")) or cluster_label(keyword, cluster_detail)
    return {
        "keyword": keyword,
        "volume": int(volume) if volume.is_integer() else volume,
        "difficulty": round(difficulty, 1),
        "intent": intent,
        "cluster": cluster,
        "url": url,
        "cpc": round(
            parse_decimal_metric(row.get("cpc") or _value(row, CPC_ALIASES)), 2
        ),
        "competition": round(
            parse_decimal_metric(
                row.get("competition") or _value(row, COMPETITION_ALIASES)
            ),
            3,
        ),
        "results": int(
            parse_metric(row.get("results") or _value(row, RESULTS_ALIASES))
        ),
        "trend": clean_text(row.get("trend") or _value(row, TREND_ALIASES)),
        "priority": clean_text(row.get("priority")) or priority_from_metrics(volume, difficulty),
        "rank": clean_text(position),
        "status": clean_text(row.get("status"))
        or ("Mapped" if url else "Chưa map Landing Page"),
        "source": clean_text(row.get("source")) or source,
    }


def priority_from_metrics(volume, difficulty):
    if volume >= 1000 and difficulty <= 65:
        return "P0"
    if volume >= 100 or difficulty <= 75:
        return "P1"
    return "P2"


def deduplicate_keywords(rows):
    merged = {}
    order = []
    for raw in rows:
        row = dict(raw)
        key = ascii_key(row.get("keyword"))
        if not key:
            continue
        if key not in merged:
            merged[key] = row
            order.append(key)
            continue
        current = merged[key]
        current["volume"] = max(
            parse_metric(current.get("volume")), parse_metric(row.get("volume"))
        )
        current["difficulty"] = max(
            parse_metric(current.get("difficulty")),
            parse_metric(row.get("difficulty")),
        )
        current["cpc"] = max(
            parse_metric(current.get("cpc")), parse_metric(row.get("cpc"))
        )
        for field in ("url", "rank", "trend", "cluster"):
            if not clean_text(current.get(field)) and clean_text(row.get(field)):
                current[field] = row[field]
        if current.get("source") != row.get("source"):
            current["source"] = "Merged"
    return [merged[key] for key in order]


def rebuild_keyword_intelligence(rows, cluster_detail="Balanced"):
    normalized = []
    for row in rows:
        item = normalize_keyword_row(
            row,
            source=clean_text(row.get("source")) or "Manual",
            cluster_detail=cluster_detail,
        )
        item["intent"] = classify_intent(item["keyword"])
        item["cluster"] = cluster_label(item["keyword"], cluster_detail)
        normalized.append(item)
    return deduplicate_keywords(normalized)


def cluster_summary(rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[clean_text(row.get("cluster")) or "Chưa phân nhóm"].append(row)
    result = []
    for cluster, members in grouped.items():
        intents = Counter(row.get("intent") for row in members)
        primary = max(
            members,
            key=lambda row: (
                parse_metric(row.get("volume")),
                -parse_metric(row.get("difficulty")),
            ),
        )
        total_volume = sum(parse_metric(row.get("volume")) for row in members)
        avg_kd = (
            sum(parse_metric(row.get("difficulty")) for row in members) / len(members)
        )
        result.append(
            {
                "cluster": cluster,
                "keywords": len(members),
                "total_volume": int(total_volume),
                "avg_difficulty": round(avg_kd, 1),
                "dominant_intent": intents.most_common(1)[0][0] if intents else "",
                "primary_keyword": primary.get("keyword", ""),
                "landing_page": next(
                    (row.get("url") for row in members if row.get("url")), ""
                ),
            }
        )
    return sorted(
        result,
        key=lambda row: (-row["total_volume"], row["cluster"].lower()),
    )


def keyword_kpis(rows):
    rows = list(rows)
    intents = Counter(row.get("intent") for row in rows)
    return {
        "keywords": len(rows),
        "total_volume": int(sum(parse_metric(row.get("volume")) for row in rows)),
        "average_difficulty": round(
            sum(parse_metric(row.get("difficulty")) for row in rows) / len(rows), 1
        )
        if rows
        else 0,
        "clusters": len({row.get("cluster") for row in rows if row.get("cluster")}),
        "mapped": sum(bool(row.get("url")) for row in rows),
        "intents": dict(intents),
    }


class SemrushKeywordImporter:
    def read(self, path, cluster_detail="Balanced"):
        path = Path(path)
        suffix = path.suffix.lower()
        if suffix in (".xlsx", ".xlsm"):
            raw_rows, headers = self._read_xlsx(path)
        elif suffix in (".csv", ".tsv", ".txt"):
            raw_rows, headers = self._read_delimited(path)
        else:
            raise ValueError("Chỉ hỗ trợ file Semrush CSV, TSV hoặc XLSX.")
        keyword_column = self._find_header(headers, KEYWORD_ALIASES)
        if not keyword_column:
            raise ValueError(
                "Không tìm thấy cột Keyword. Hãy export lại Semrush với cột Keyword."
            )
        rows = [
            normalize_keyword_row(
                raw,
                source="Semrush",
                cluster_detail=cluster_detail,
            )
            for raw in raw_rows
            if clean_text(_value(raw, KEYWORD_ALIASES))
        ]
        before = len(rows)
        rows = deduplicate_keywords(rows)
        return rows, {
            "file": path.name,
            "rows_read": before,
            "keywords_imported": len(rows),
            "duplicates_removed": before - len(rows),
            "headers": headers,
            "detected": {
                "keyword": keyword_column,
                "volume": self._find_header(headers, VOLUME_ALIASES),
                "difficulty": self._find_header(headers, DIFFICULTY_ALIASES),
                "intent": self._find_header(headers, INTENT_ALIASES),
                "cpc": self._find_header(headers, CPC_ALIASES),
                "url": self._find_header(headers, URL_ALIASES),
            },
        }

    @staticmethod
    def _find_header(headers, aliases):
        by_key = {ascii_key(header): header for header in headers}
        for alias in aliases:
            if ascii_key(alias) in by_key:
                return by_key[ascii_key(alias)]
        return ""

    @staticmethod
    def _read_delimited(path):
        encodings = ("utf-8-sig", "utf-16", "cp1258", "latin-1")
        last_error = None
        for encoding in encodings:
            try:
                with open(path, "r", encoding=encoding, newline="") as handle:
                    sample = handle.read(8192)
                    handle.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
                    except csv.Error:
                        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
                    reader = csv.DictReader(handle, dialect=dialect)
                    rows = [
                        {
                            clean_text(key): clean_text(value)
                            for key, value in row.items()
                            if key is not None
                        }
                        for row in reader
                    ]
                    return rows, list(reader.fieldnames or [])
            except UnicodeError as exc:
                last_error = exc
        raise ValueError(f"Không đọc được encoding của file: {last_error}")

    @staticmethod
    def _read_xlsx(path):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Thiếu openpyxl để đọc file Semrush XLSX.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [clean_text(value) for value in next(iterator)]
        except StopIteration:
            return [], []
        rows = []
        for values in iterator:
            rows.append(
                {
                    headers[index]: clean_text(value)
                    for index, value in enumerate(values)
                    if index < len(headers)
                }
            )
        workbook.close()
        return rows, headers


def export_keyword_workbook(path, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError("Thiếu openpyxl để xuất Excel.") from exc

    rows = list(rows)
    summaries = cluster_summary(rows)
    kpis = keyword_kpis(rows)
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Tổng quan"
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:F2")
    dashboard["A1"] = "SEO AI STUDIO V4.4 - KEYWORD INTELLIGENCE"
    dashboard["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    dashboard["A1"].fill = PatternFill("solid", fgColor="17324D")
    dashboard["A1"].alignment = Alignment(horizontal="left", vertical="center")
    metrics = [
        ("Keywords", kpis["keywords"]),
        ("Total Volume", kpis["total_volume"]),
        ("Average KD", kpis["average_difficulty"]),
        ("Clusters", kpis["clusters"]),
        ("Mapped URL", kpis["mapped"]),
    ]
    for index, (label, value) in enumerate(metrics, 1):
        dashboard.cell(4, index, label)
        dashboard.cell(5, index, value)
        dashboard.cell(4, index).font = Font(bold=True, color="1769C2")
        dashboard.cell(4, index).fill = PatternFill("solid", fgColor="E9F3FF")
        dashboard.cell(5, index).font = Font(size=16, bold=True, color="17324D")
        dashboard.cell(4, index).alignment = Alignment(horizontal="center")
        dashboard.cell(5, index).alignment = Alignment(horizontal="center")
        dashboard.column_dimensions[chr(64 + index)].width = 20
    dashboard["A8"] = "Search Intent"
    dashboard["A8"].font = Font(size=13, bold=True, color="17324D")
    for offset, intent in enumerate(
        ("Informational", "Commercial", "Transactional", "Navigational"), 9
    ):
        dashboard.cell(offset, 1, intent)
        dashboard.cell(offset, 2, kpis["intents"].get(intent, 0))
        dashboard.cell(offset, 1).fill = PatternFill(
            "solid", fgColor=INTENT_COLORS[intent].lstrip("#")
        )

    thin = Side(style="thin", color="DDE3EA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def make_sheet(name, columns, data):
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        for col, (_, label, width) in enumerate(columns, 1):
            cell = sheet.cell(1, col, label)
            cell.fill = PatternFill("solid", fgColor="1769C2")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center")
            sheet.column_dimensions[cell.column_letter].width = width
        for row_index, source in enumerate(data, 2):
            for col, (key, _, _) in enumerate(columns, 1):
                cell = sheet.cell(row_index, col, source.get(key, ""))
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        if data:
            name_key = re.sub(r"[^A-Za-z0-9]", "", name) + "Table"
            table = Table(displayName=name_key, ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True
            )
            sheet.add_table(table)
        return sheet

    keyword_sheet = make_sheet(
        "Keywords",
        [
            ("keyword", "Keyword", 40),
            ("volume", "Volume", 14),
            ("difficulty", "KD %", 12),
            ("intent", "Search Intent", 18),
            ("cluster", "Cluster", 28),
            ("url", "Landing Page", 50),
            ("cpc", "CPC", 12),
            ("rank", "Rank", 10),
            ("priority", "Priority", 11),
            ("status", "Status", 22),
            ("source", "Source", 14),
        ],
        rows,
    )
    if rows:
        keyword_sheet.conditional_formatting.add(
            f"C2:C{len(rows) + 1}",
            CellIsRule(
                operator="greaterThanOrEqual",
                formula=["70"],
                fill=PatternFill("solid", fgColor="F8C7C7"),
            ),
        )
    make_sheet(
        "Clusters",
        [
            ("cluster", "Cluster", 30),
            ("keywords", "Keywords", 12),
            ("total_volume", "Total Volume", 16),
            ("avg_difficulty", "Avg. KD", 12),
            ("dominant_intent", "Dominant Intent", 20),
            ("primary_keyword", "Primary Keyword", 38),
            ("landing_page", "Landing Page", 50),
        ],
        summaries,
    )
    dashboard["A15"] = "Tạo bởi SEO AI Studio V4.4"
    dashboard["B15"] = datetime.now().astimezone().isoformat(timespec="seconds")
    workbook.calculation.fullCalcOnLoad = True
    workbook.save(path)
    return Path(path)
